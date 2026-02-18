import json
import os
import re
import shutil
import time
from datetime import datetime

from openpyxl import load_workbook


def load_settings() -> dict:
    with open("settings.json", "r", encoding="utf-8") as f:
        return json.load(f)


SET = load_settings()

DB_PATH = SET["database_path"]
EXCEL_PATH = SET["excel_path"]
BACKUP_DIR = SET["backup_folder"]

FOLIO_RE = re.compile(r"^RMA-RED-(\d{4})$")


def _read_db() -> list:
    if not os.path.exists(DB_PATH):
        return []
    with open(DB_PATH, "r", encoding="utf-8") as f:
        try:
            data = json.load(f)
            return data if isinstance(data, list) else []
        except json.JSONDecodeError:
            return []


def _write_db(data: list) -> None:
    with open(DB_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def backup_db() -> str | None:
    if not os.path.exists(DB_PATH):
        return None
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    dst = os.path.join(BACKUP_DIR, f"rma_database_{ts}.json")
    shutil.copyfile(DB_PATH, dst)
    return dst


def highest_folio_in_db() -> int:
    maxn = 0
    for rec in _read_db():
        fol = rec.get("folio") if isinstance(rec, dict) else None
        if not fol:
            continue
        match = FOLIO_RE.match(fol)
        if match:
            maxn = max(maxn, int(match.group(1)))
    return maxn


def highest_folio_in_excel() -> int:
    if not os.path.exists(EXCEL_PATH):
        return 0
    try:
        wb = load_workbook(EXCEL_PATH, data_only=True)
        ws = wb.active
        maxn = 0
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 1:
                continue
            fol = (str(row[0]).strip() if row and len(row) > 0 and row[0] else "")
            match = FOLIO_RE.match(fol)
            if match:
                maxn = max(maxn, int(match.group(1)))
        return maxn
    except OSError:
        return 0


def get_next_folio() -> str:
    nxt = max(highest_folio_in_db(), highest_folio_in_excel()) + 1
    return f"RMA-RED-{nxt:04d}"


def get_all() -> list:
    return _read_db()


def get_by_folio(folio: str):
    for rec in _read_db():
        if isinstance(rec, dict) and rec.get("folio") == folio:
            return rec
    return None


def upsert_entry(entry: dict) -> None:
    if "folio" not in entry:
        raise ValueError("El registro no tiene folio.")

    data = _read_db()
    idx = next(
        (i for i, rec in enumerate(data) if isinstance(rec, dict) and rec.get("folio") == entry["folio"]),
        None,
    )
    if idx is None:
        data.append(entry)
    else:
        data[idx] = entry
    _write_db(data)


def save_rma_data(entry: dict) -> None:
    required = [
        "folio",
        "fecha_recepcion",
        "nombre_cliente",
        "factura",
        "producto",
        "marca",
        "modelo",
        "serie",
        "falla",
        "responsable",
    ]
    for key in required:
        if key not in entry:
            raise ValueError(f"Falta el campo requerido: {key}")
    upsert_entry(entry)


def excel_headers() -> list:
    return SET["excel_headers"]


def write_or_update_excel_row(entry: dict) -> None:
    headers = excel_headers()
    if not os.path.exists(EXCEL_PATH):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "RMA"
        ws.append(headers)
        wb.save(EXCEL_PATH)

    wb = load_workbook(EXCEL_PATH)
    ws = wb.active

    col_map = {headers[i]: i + 1 for i in range(len(headers))}
    folio_col = col_map.get("Folio", 1)

    target_row = None
    folio = str(entry.get("folio", ""))
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if (str(row[folio_col - 1]).strip() if row[folio_col - 1] else "") == folio:
            target_row = i
            break

    if target_row is None:
        target_row = ws.max_row + 1

    values = [
        entry.get("folio", ""),
        entry.get("fecha_recepcion", ""),
        entry.get("nombre_cliente", ""),
        entry.get("telefono", ""),
        entry.get("factura", ""),
        entry.get("producto", ""),
        entry.get("marca", ""),
        entry.get("modelo", ""),
        entry.get("serie", ""),
        entry.get("falla", ""),
        entry.get("observaciones", ""),
        entry.get("responsable", ""),
        entry.get("estado", "En Proceso"),
        entry.get("fecha_entrega", ""),
        entry.get("tipo_resolucion", ""),
        entry.get("producto_nuevo", ""),
        entry.get("modelo_nuevo", ""),
        entry.get("serie_nueva", ""),
    ]

    for c, value in enumerate(values, start=1):
        ws.cell(row=target_row, column=c, value=value)

    from openpyxl.styles import Font, PatternFill

    font = Font(name="Arial", size=12)
    estado_idx = headers.index("Estado") + 1
    for c in range(1, len(headers) + 1):
        ws.cell(row=target_row, column=c).font = font

    estado = (entry.get("estado") or "").strip().lower()
    fill = None
    if estado == "en proceso":
        fill = PatternFill("solid", fgColor="FFF8B3")
    elif estado in ["aprobada", "válida", "valida"]:
        fill = PatternFill("solid", fgColor="C6EFCE")
    elif estado in ["no válida", "no valida"]:
        fill = PatternFill("solid", fgColor="FFC7CE")

    if fill:
        ws.cell(row=target_row, column=estado_idx).fill = fill

    for _ in range(5):
        try:
            wb.save(EXCEL_PATH)
            return
        except PermissionError:
            time.sleep(1)
    raise PermissionError(
        f"No se pudo guardar el archivo '{EXCEL_PATH}' después de 5 intentos."
    )


def _to_str(v) -> str:
    return "" if v is None else str(v).strip()


def import_from_excel_to_db(deduplicate: bool = True, backup: bool = True) -> int:
    if not os.path.exists(EXCEL_PATH):
        return 0
    if backup:
        backup_db()

    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    hm = {h: i for i, h in enumerate(headers)}

    required_headers = [
        "Folio",
        "Fecha Recepción",
        "Cliente",
        "Factura",
        "Producto",
        "Marca",
        "Modelo",
        "Serie",
        "Falla",
        "Responsable",
        "Estado",
    ]
    missing_headers = [h for h in required_headers if h not in hm]
    if missing_headers:
        raise ValueError(f"Faltan columnas obligatorias en Excel: {', '.join(missing_headers)}")

    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        folio = _to_str(row[hm["Folio"]])
        if not folio:
            continue

        rec = {
            "folio": folio,
            "fecha_recepcion": _to_str(row[hm.get("Fecha Recepción", 1)]),
            "nombre_cliente": _to_str(row[hm.get("Cliente", 2)]),
            "telefono": _to_str(row[hm.get("Teléfono", 3)]),
            "factura": _to_str(row[hm.get("Factura", 4)]),
            "producto": _to_str(row[hm.get("Producto", 5)]),
            "marca": _to_str(row[hm.get("Marca", 6)]),
            "modelo": _to_str(row[hm.get("Modelo", 7)]),
            "serie": _to_str(row[hm.get("Serie", 8)]),
            "falla": _to_str(row[hm.get("Falla", 9)]),
            "observaciones": _to_str(row[hm.get("Observaciones", 10)]),
            "responsable": _to_str(row[hm.get("Responsable", 11)]),
            "estado": _to_str(row[hm.get("Estado", 12)]) or "En Proceso",
            "fecha_entrega": _to_str(row[hm.get("Fecha de Resolución", 13)]),
            "tipo_resolucion": _to_str(row[hm.get("Tipo de Resolución", 14)]),
            "producto_nuevo": _to_str(row[hm.get("Producto Nuevo", 15)]),
            "modelo_nuevo": _to_str(row[hm.get("Modelo Nuevo", 16)]),
            "serie_nueva": _to_str(row[hm.get("Serie Nueva", 17)]),
        }
        upsert_entry(rec)
        count += 1

    return count
